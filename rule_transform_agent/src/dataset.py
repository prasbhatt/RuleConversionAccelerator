"""
Loads the ground-truth Excel workbook produced during the data preparation
phase (the "Synthetic Rules" sheet) and exposes it as a list of plain
RuleExample objects, plus a helper to export it to JSONL for embedding.
"""
from __future__ import annotations
from dataclasses import dataclass, asdict
import json
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class RuleExample:
    rule_id: str
    source_rule_id: str
    source_rule_name: str
    source_rule_type: str
    source_rule_definition: str
    rule_category: str
    complexity: str
    target_attribute: str | None
    target_locale: str | None
    referenced_attributes: str
    expected_target_json: str
    dataset_split: str

    def embedding_text(self) -> str:
        """
        The text that gets embedded for retrieval. It deliberately includes
        the category and complexity as plain words, because that lets a
        similarity search naturally cluster same-category / same-complexity
        rules even before any metadata filter is applied.
        """
        return (
            f"Category: {self.rule_category}\n"
            f"Complexity: {self.complexity}\n"
            f"Source rule name: {self.source_rule_name}\n"
            f"Source rule definition:\n{self.source_rule_definition}"
        )


def load_ground_truth(xlsx_path: Path, sheet_name: str = "Synthetic Rules") -> list[RuleExample]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {h: i for i, h in enumerate(headers)}

    def col(row, name, default=None):
        i = index.get(name)
        return row[i] if i is not None else default

    examples: list[RuleExample] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not col(row, "Synthetic Rule ID"):
            continue
        examples.append(RuleExample(
            rule_id=str(col(row, "Synthetic Rule ID")),
            source_rule_id=str(col(row, "Source Rule ID")),
            source_rule_name=str(col(row, "Source Rule Name")),
            source_rule_type=str(col(row, "Source Rule Type") or ""),
            source_rule_definition=str(col(row, "Source Rule Definition") or ""),
            rule_category=str(col(row, "Rule Category") or col(row, "Rule Type") or ""),
            complexity=str(col(row, "Complexity") or ""),
            target_attribute=col(row, "Target Attribute"),
            target_locale=col(row, "Target Locale"),
            referenced_attributes=str(col(row, "Referenced Attributes") or ""),
            expected_target_json=str(col(row, "Expected Target JSON") or ""),
            dataset_split=str(col(row, "Dataset Split") or ""),
        ))
    return examples


def export_jsonl(examples: list[RuleExample], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as f:
        for ex in examples:
            f.write(json.dumps(asdict(ex), ensure_ascii=False) + "\n")


def filter_split(examples: list[RuleExample], split: str) -> list[RuleExample]:
    return [e for e in examples if e.dataset_split.lower() == split.lower()]
